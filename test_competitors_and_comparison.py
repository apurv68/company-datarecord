"""
Unit Test Suite for Competitor Benchmarking (Table #6) & Head-to-Head Comparison Mode
=====================================================================================
Covers:
1. Regex parsing for dual-company queries ("Company A vs Company B", "A versus B", "A vs. B")
2. Comparative delta calculation (calculate_comparative_edge)
3. Table #6 Peer extraction and schema compliance
4. Dual-company Word document generation (save_comparison_docx)
5. Single-company Word document with Section 6.0 (save_company_docx)
6. Table #6 CSV export (company_peers.csv)
"""

import os
import re
import unittest
import docx
from company_lookup import (
    calculate_comparative_edge,
    save_comparison_docx,
    save_company_docx,
    save_table_records,
    extract_office_address,
    EXPORT_TABLE6_CSV_PATH
)


class TestCompetitorsAndComparison(unittest.TestCase):

    def test_head_to_head_query_regex(self):
        """Verify regex accurately splits various 'vs' formats without false positives."""
        pattern = r"\s+(?:vs\.?|versus)\s+"

        # Standard 'vs'
        parts1 = re.split(pattern, "TCS vs Infosys", flags=re.I)
        self.assertEqual(len(parts1), 2)
        self.assertEqual(parts1[0].strip(), "TCS")
        self.assertEqual(parts1[1].strip(), "Infosys")

        # 'versus'
        parts2 = re.split(pattern, "Liberty Shoes versus Bata India", flags=re.I)
        self.assertEqual(len(parts2), 2)
        self.assertEqual(parts2[0].strip(), "Liberty Shoes")
        self.assertEqual(parts2[1].strip(), "Bata India")

        # 'vs.' with period
        parts3 = re.split(pattern, "Zomato vs. Swiggy", flags=re.I)
        self.assertEqual(len(parts3), 2)
        self.assertEqual(parts3[0].strip(), "Zomato")
        self.assertEqual(parts3[1].strip(), "Swiggy")

        # Single company must NOT split
        parts_single = re.split(pattern, "Liberty Shoes Limited", flags=re.I)
        self.assertEqual(len(parts_single), 1)

    def test_calculate_comparative_edge(self):
        """Test delta and multiplier computation between two companies."""
        # Multiplier case (Bata is >2x Liberty Shoes)
        edge1 = calculate_comparative_edge("Current Market Cap", "Rs 374 Cr", "Rs 8135 Cr", "Liberty Shoes", "Bata India")
        self.assertIn("Bata", edge1)
        self.assertIn("x larger", edge1)

        # Percentage lead case
        edge2 = calculate_comparative_edge("Latest Net Revenue", "Rs 1500 Cr", "Rs 1000 Cr", "Company Alpha", "Company Beta")
        self.assertIn("Alpha", edge2)
        self.assertIn("+50%", edge2)

        # Parity case
        edge3 = calculate_comparative_edge("Employee Headcount", "5,000", "5,000", "Alpha", "Beta")
        self.assertEqual(edge3, "Parity (Equal)")

        # N/A case
        edge4 = calculate_comparative_edge("Share Price", "N/A (Privately Held)", "Rs 450", "Alpha", "Beta")
        self.assertEqual(edge4, "N/A")

    def test_save_comparison_docx_creates_valid_document(self):
        """Verify dedicated Head-to-Head Word document generation."""
        comp_a = {
            "Company Name": "Tata Consultancy Services Ltd",
            "Stock Ticker": "NSE/BSE: TCS",
            "Business Type": "Public Limited",
            "Headquarter (City)": "Mumbai, Maharashtra",
            "CEO": "K. Krithivasan",
            "CFO": "Samir Seksaria",
            "Market Cap": "₹ 764,501 Cr",
            "Share Price": "₹ 2,113.00",
            "Net Revenue": "Rs 240,893 Cr",
            "Net Profit": "Rs 46,099 Cr",
            "EBITDA": "Rs 62,500 Cr",
            "Headcount": "601,546",
            "Growth Verdict": "Steady Growth"
        }

        comp_b = {
            "Company Name": "Infosys Limited",
            "Stock Ticker": "NSE/BSE: INFY",
            "Business Type": "Public Limited",
            "Headquarter (City)": "Bengaluru, Karnataka",
            "CEO": "Salil Parekh",
            "CFO": "Jayesh Sanghrajka",
            "Market Cap": "₹ 416,054 Cr",
            "Share Price": "₹ 1,025.20",
            "Net Revenue": "Rs 153,670 Cr",
            "Net Profit": "Rs 26,248 Cr",
            "EBITDA": "Rs 38,400 Cr",
            "Headcount": "317,240",
            "Growth Verdict": "Steady Growth"
        }

        test_out = "scratch_test_tcs_vs_infy.docx"
        try:
            saved_path = save_comparison_docx(comp_a, comp_b, output_filepath=test_out)
            self.assertIsNotNone(saved_path)
            self.assertTrue(os.path.exists(test_out))
            self.assertGreater(os.path.getsize(test_out), 1000)

            # Validate docx content
            doc = docx.Document(test_out)
            text_blob = "\n".join([p.text for p in doc.paragraphs] + [c.text for tbl in doc.tables for r in tbl.rows for c in r.cells])
            self.assertIn("Tata Consultancy Services Ltd vs Infosys Limited", text_blob)
            self.assertIn("HEAD-TO-HEAD INSTITUTIONAL COMPETITIVE INTELLIGENCE DOSSIER", text_blob)
            self.assertIn("1.0 Corporate Dimensions & Leadership Showdown", text_blob)
            self.assertIn("K. Krithivasan", text_blob)
            self.assertIn("Salil Parekh", text_blob)
            self.assertIn("TCS (+84%)", text_blob)  # Market Cap lead
        finally:
            if os.path.exists(test_out):
                try:
                    os.remove(test_out)
                except Exception:
                    pass

    def test_save_company_docx_with_table6_section(self):
        """Verify Table #6 Peer Comparison matrix is correctly embedded into single-company DOCX."""
        data1 = {
            "Company Name": "Liberty Shoes Limited",
            "Business Type": "Public Limited",
            "Is Listed Company": "Yes",
            "Stock Ticker": "NSE/BSE: LIBERTSHOE",
            "Managing Director / CEO": "Adesh Kumar Gupta"
        }
        sources1 = [{"name": "BSE India", "url": "https://www.bseindia.com"}]

        data2 = {
            "Company Name": "Liberty Shoes Limited",
            "periods": ["Mar 2024"],
            "rows": [{"Fiscal Period / Year": "Mar 2024", "Market Cap": "Rs 374 Cr", "Net Revenue/Net Sales": "Rs 650 Cr", "Net Profit": "Rs 14 Cr", "EBITDA": "Rs 45 Cr", "Employee Headcount": "2,100"}]
        }
        sources2 = [{"name": "Screener.in", "url": "https://www.screener.in"}]

        data6 = {
            "Target Company": "Liberty Shoes Limited",
            "rows": [
                {"s_no": "1.", "name": "Bata India", "cmp": "632.95", "pe": "46.80", "market_cap": "8135.13", "roce": "12.74", "sales_qtr": "978.95", "qtr_profit_var": "15.88", "is_benchmark": False, "is_target": False},
                {"s_no": "2.", "name": "Liberty Shoes", "cmp": "219.60", "pe": "43.92", "market_cap": "374.19", "roce": "7.91", "sales_qtr": "169.86", "qtr_profit_var": "-76.97", "is_benchmark": False, "is_target": True},
                {"s_no": "", "name": "Median: 10 Co.", "cmp": "215.65", "pe": "40.73", "market_cap": "3544.49", "roce": "15.41", "sales_qtr": "277.53", "qtr_profit_var": "13.79", "is_benchmark": True, "is_target": False}
            ]
        }
        sources6 = [{"name": "Screener.in Peer Benchmarks", "url": "https://www.screener.in/api/company/3178/peers/"}]

        test_out = "scratch_test_liberty_t6.docx"
        try:
            saved = save_company_docx(
                data1, sources1, data2, sources2,
                data6=data6, sources6=sources6,
                output_filepath=test_out
            )
            self.assertIsNotNone(saved)
            self.assertTrue(os.path.exists(test_out))

            doc = docx.Document(test_out)
            text_blob = "\n".join([p.text for p in doc.paragraphs] + [c.text for tbl in doc.tables for r in tbl.rows for c in r.cells])
            self.assertIn("6.0 Peer Comparison & Competitive Benchmarking Matrix", text_blob)
            self.assertIn("Bata India", text_blob)
            self.assertIn("▶ Liberty Shoes", text_blob)
            self.assertIn("Median: 10 Co.", text_blob)
        finally:
            if os.path.exists(test_out):
                try:
                    os.remove(test_out)
                except Exception:
                    pass

    def test_save_table_records_writes_company_peers_csv(self):
        """Verify Table #6 writes to company_peers.csv cleanly."""
        data1 = {"Company Name": "Liberty Shoes Limited"}
        data2 = {"rows": []}
        data6 = {
            "Target Company": "Liberty Shoes Limited",
            "rows": [
                {"name": "Bata India", "cmp": "632.95", "pe": "46.80", "market_cap": "8135.13", "div_yield": "3.94", "np_qtr": "63.98", "qtr_profit_var": "15.88", "sales_qtr": "978.95", "qtr_sales_var": "3.94", "roce": "12.74", "is_benchmark": False, "is_target": False}
            ]
        }
        sources6 = [{"name": "Screener", "url": "https://www.screener.in"}]

        test_csv1 = "scratch_test_records.csv"
        test_csv2 = "scratch_test_fin.csv"
        test_csv6 = "scratch_test_peers.csv"
        test_json = "scratch_test_peers.json"
        try:
            save_table_records(
                data1, [], data2, [],
                data6=data6, sources6=sources6,
                csv1_path=test_csv1,
                csv2_path=test_csv2,
                csv6_path=test_csv6,
                json_path=test_json
            )
            self.assertTrue(os.path.exists(test_csv6))
            with open(test_csv6, "r", encoding="utf-8") as f:
                content = f.read()
            self.assertIn("Target Company,Peer Company,CMP (Rs)", content)
            self.assertIn("Liberty Shoes Limited", content)
        finally:
            for p in [test_csv1, test_csv2, test_csv6, test_json]:
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except Exception:
                        pass

    def test_extract_office_address_strips_preamble(self):
        """Verify extract_office_address removes sentence preambles, verb prefixes, and entity names."""
        raw_agreeya = "AGREEYA SOLUTIONS (INDIA) PVT LTD's registered office address is D-71 Amar Colony Lajpat Nagar Iv, South Delhi, New Delhi, Delhi, India, 110024"
        addr1 = extract_office_address(raw_agreeya, company_name="AgreeYa Solutions")
        self.assertNotIn("AGREEYA", addr1)
        self.assertNotIn("registered office address is", addr1)
        self.assertIn("D-71 Amar Colony", addr1)
        self.assertIn("110024", addr1)

        raw_woodofa = "Contact Information Get in Touch Head Office A-3, WHS, Timber Market, Kirti Nagar, Delhi-110015 +91-9315950993 sales@woodofa.com"
        addr2 = extract_office_address(raw_woodofa, company_name="Woodofa")
        self.assertNotIn("Contact Information", addr2)
        self.assertNotIn("Get in Touch", addr2)
        self.assertNotIn("Head Office", addr2)
        self.assertIn("A-3, WHS, Timber Market", addr2)
        self.assertIn("110015", addr2)

    def test_save_company_docx_table1_deduplication(self):
        """Verify Table 1 does not generate duplicate rows for CEO, Address, or Business Type."""
        data1 = {
            "Company Name": "Woodofa",
            "CEO": "Tushar Mittal",
            "Office Address": "A-3, WHS, Timber Market, Kirti Nagar, Delhi-110015",
            "Business Type (Private Limited/Public Limited)": "Private Limited",
            "Headquarter (City)": "Noida",
            "Founding Year": "2013"
        }
        test_out = "scratch_test_dedup.docx"
        try:
            saved = save_company_docx(data1, [], {"rows": []}, [], output_filepath=test_out)
            self.assertIsNotNone(saved)
            doc = docx.Document(test_out)
            t1_rows = [r.cells[0].text.strip() for r in doc.tables[0].rows]
            # Ensure each dimension appears exactly once
            self.assertEqual(t1_rows.count("Managing Director / CEO"), 1)
            self.assertEqual(t1_rows.count("CEO"), 0)
            self.assertEqual(t1_rows.count("Office Address"), 1)
            self.assertEqual(t1_rows.count("Registered Address"), 0)
            self.assertEqual(t1_rows.count("Business Type"), 1)
            self.assertEqual(t1_rows.count("Business Type (Private Limited/Public Limited)"), 0)
        finally:
            if os.path.exists(test_out):
                try:
                    os.remove(test_out)
                except Exception:
                    pass


if __name__ == "__main__":
    unittest.main()
